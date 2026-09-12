#!/usr/bin/env python3
"""Generate customer brief docx from BRIEF.md."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[3]
BRIEF_MD = ROOT / "project-docs" / "Common" / "actual" / "BRIEF.md"
OUTPUT = ROOT / "project-docs" / "Common" / "actual" / "BRIEF-zakazchika.docx"


def strip_md_bold(text: str) -> tuple[str, list[tuple[int, int]]]:
    """Return plain text and bold spans (start, end) in plain text."""
    plain = ""
    bold_spans: list[tuple[int, int]] = []
    i = 0
    while i < len(text):
        if text[i : i + 2] == "**":
            end = text.find("**", i + 2)
            if end == -1:
                plain += text[i:]
                break
            start = len(plain)
            segment = text[i + 2 : end]
            plain += segment
            bold_spans.append((start, len(plain)))
            i = end + 2
        else:
            plain += text[i]
            i += 1
    return plain, bold_spans


def add_rich_paragraph(cell, text: str, size: Pt = Pt(10)) -> None:
    p = cell.paragraphs[0]
    p.clear()
    plain, spans = strip_md_bold(text)
    if not spans:
        run = p.add_run(plain)
        run.font.size = size
        run.font.name = "Calibri"
        return
    pos = 0
    bold_set = set()
    for start, end in spans:
        bold_set.update(range(start, end))
    while pos < len(plain):
        next_bold_change = len(plain)
        is_bold = pos in bold_set
        while next_bold_change > pos and (next_bold_change - 1 in bold_set) == is_bold:
            next_bold_change -= 1
        j = pos + 1
        while j < len(plain) and (j in bold_set) == is_bold:
            j += 1
        run = p.add_run(plain[pos:j])
        run.font.size = size
        run.font.name = "Calibri"
        run.bold = is_bold
        pos = j


def set_cell_shading(cell, fill: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        from docx.oxml import OxmlElement

        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], col_widths_cm: list[float]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0].cells
    for idx, title in enumerate(headers):
        hdr[idx].text = ""
        add_rich_paragraph(hdr[idx], title, Pt(10))
        for run in hdr[idx].paragraphs[0].runs:
            run.bold = True
        set_cell_shading(hdr[idx], "E8EEF7")

    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            add_rich_paragraph(cells[c_idx], val, Pt(10))
            if c_idx == len(row) - 1 and headers[-1].startswith("Ответ"):
                cells[c_idx].paragraphs[0].add_run("\n")

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)

    # column widths via table layout fixed
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        from docx.oxml import OxmlElement

        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        from docx.oxml import OxmlElement

        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl_pr.append(layout)

    for idx, width in enumerate(col_widths_cm):
        for row in table.rows:
            row.cells[idx].width = Cm(width)


def parse_brief(md_text: str) -> dict:
    sections: list[dict] = []
    current: dict | None = None
    parent_section: dict | None = None
    in_table = False
    table_headers: list[str] = []
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal in_table, table_headers, table_rows
        if current is not None and in_table and table_headers:
            current.setdefault("tables", []).append(
                {"headers": table_headers, "rows": table_rows}
            )
        in_table = False
        table_headers = []
        table_rows = []

    for line in md_text.splitlines():
        if line.startswith("## "):
            flush_table()
            title = line[3:].strip()
            current = {"title": title, "intro": [], "tables": [], "subsections": []}
            parent_section = current
            sections.append(current)
            continue

        if line.startswith("### "):
            flush_table()
            if parent_section is not None:
                sub = {"title": line[4:].strip(), "tables": []}
                parent_section.setdefault("subsections", []).append(sub)
                current = sub
            continue

        if line.strip().startswith("|") and "---" not in line:
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not in_table:
                in_table = True
                table_headers = cells
                table_rows = []
            else:
                table_rows.append(cells)
            continue

        if in_table and not line.strip().startswith("|"):
            flush_table()

        if line.startswith("> ") and current is not None:
            current["intro"].append(line[2:].strip())
        elif line.strip() == "---":
            flush_table()
        elif (
            line.strip()
            and not line.startswith("#")
            and current is not None
            and not line.startswith("|")
            and not line.startswith("**Формат")
            and not line.startswith("**Назначение")
        ):
            if "intro" in current:
                current["intro"].append(line.strip())

    flush_table()
    return {"sections": sections}


def col_widths_for_headers(headers: list[str]) -> list[float]:
    if len(headers) == 4:
        return [1.2, 5.5, 4.0, 4.8]
    if len(headers) == 2:
        return [6.0, 9.5]
    return [1.2, 8.3, 6.0]


def build_doc(data: dict) -> Document:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Бриф заказчика")
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = "Calibri"

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("Система планирования смен и учёта ФОТ — производство")
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.add_run(
        "Пожалуйста, заполните поля «Ответ заказчика». "
        "Это нужно для точной оценки сроков и стоимости."
    )

    note = doc.add_paragraph()
    nrun = note.add_run(
        "Блок G — совместный чеклист: помогает выбрать формат работы "
        "(лёгкий MVP или инженерная разработка с запасом на рост)."
    )
    nrun.italic = True
    nrun.font.size = Pt(10)

    doc.add_paragraph()

    for section in data["sections"]:
        h = doc.add_heading(section["title"], level=1)
        h.runs[0].font.name = "Calibri"

        for para in section.get("intro", []):
            p = doc.add_paragraph()
            add_rich_paragraph_to_doc(p, para)

        for table in section.get("tables", []):
            headers = table["headers"]
            rows = table["rows"]
            add_table(doc, headers, rows, col_widths_for_headers(headers))
            doc.add_paragraph()

        for sub in section.get("subsections", []):
            sh = doc.add_heading(sub["title"], level=2)
            sh.runs[0].font.name = "Calibri"
            for table in sub.get("tables", []):
                headers = table["headers"]
                rows = table["rows"]
                add_table(doc, headers, rows, col_widths_for_headers(headers))
                doc.add_paragraph()

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("— Конец брифа —")
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    return doc


def add_rich_paragraph_to_doc(p, text: str) -> None:
    plain, spans = strip_md_bold(text)
    if not spans:
        p.add_run(plain)
        return
    bold_set = set()
    for start, end in spans:
        bold_set.update(range(start, end))
    pos = 0
    while pos < len(plain):
        j = pos + 1
        is_bold = pos in bold_set
        while j < len(plain) and (j in bold_set) == is_bold:
            j += 1
        run = p.add_run(plain[pos:j])
        run.bold = is_bold
        pos = j


def main() -> None:
    md_text = BRIEF_MD.read_text(encoding="utf-8")
    data = parse_brief(md_text)
    doc = build_doc(data)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Written: {OUTPUT}")


if __name__ == "__main__":
    main()
